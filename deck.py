"""
Deck-building helpers: image normalization, Claude plan parsing, and .pptx rendering.
Kept separate from bot.py so it can be tested without Discord.
"""

import io
import json
import re

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)   # 16:9
MARGIN = Inches(0.5)

BG = RGBColor(0xFA, 0xFA, 0xF7)
INK = RGBColor(0x1F, 0x23, 0x2B)
MUTED = RGBColor(0x5B, 0x63, 0x70)
ACCENT = RGBColor(0x3B, 0x5B, 0xDB)


# ---------------------------------------------------------------- images

def normalize_image(data: bytes, max_side: int, quality: int):
    """Return (jpeg_bytes, width, height), or None if the file isn't a readable image.
    Handles EXIF rotation, transparency, palette and grayscale modes."""
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
        im = ImageOps.exif_transpose(im)
    except Exception:
        return None

    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[-1])
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")

    im.thumbnail((max_side, max_side), Image.LANCZOS)
    out = io.BytesIO()
    im.save(out, "JPEG", quality=quality, optimize=True)
    return out.getvalue(), im.width, im.height


# ---------------------------------------------------------------- plan

PLAN_SCHEMA = """{
  "title": "deck title",
  "subtitle": "one-line subtitle",
  "slides": [
    {"title": "slide title", "caption": "short caption", "images": [0, 3], "notes": "speaker notes"}
  ]
}"""


def planning_prompt(n_images: int, topic: str | None) -> str:
    topic_line = f"Topic: {topic}" if topic else "Topic: not specified. Infer it from the images."
    return f"""You are building a slide deck from {n_images} images collected in a Discord channel.
{topic_line}
Each image above is labeled with its index.

Return ONLY a JSON object in this shape, with no prose and no code fences:
{PLAN_SCHEMA}

Rules:
- Order slides so the deck tells a coherent story.
- Each slide uses 1-4 images. Prefer 1; group images only when they belong together.
- Use each image at most once. You may skip near-duplicates, blurry shots, or off-topic images.
- Slide titles: 8 words or fewer. Captions: 25 words or fewer.
- Describe what is actually visible; do not invent names, dates, or facts.
- Notes: 2-4 sentences a presenter could say aloud."""


def parse_plan(text: str, n_images: int) -> dict:
    """Extract and sanitize Claude's JSON plan. Falls back to one image per slide."""
    plan = {}
    m = re.search(r"\{.*\}", text, re.S)
    if m:
        try:
            plan = json.loads(m.group(0))
        except json.JSONDecodeError:
            plan = {}

    slides, used = [], set()
    for s in plan.get("slides", []) or []:
        if not isinstance(s, dict):
            continue
        idx = []
        for i in s.get("images", []) or []:
            if isinstance(i, int) and 0 <= i < n_images and i not in used:
                idx.append(i)
                used.add(i)
        idx = idx[:4]
        if not idx:
            continue
        slides.append({
            "title": str(s.get("title", "")).strip(),
            "caption": str(s.get("caption", "")).strip(),
            "images": idx,
            "notes": str(s.get("notes", "")).strip(),
        })

    if not slides:
        slides = [{"title": f"Image {i + 1}", "caption": "", "images": [i], "notes": ""}
                  for i in range(n_images)]

    return {
        "title": str(plan.get("title") or "Collected Images").strip(),
        "subtitle": str(plan.get("subtitle") or "").strip(),
        "slides": slides,
    }


# ---------------------------------------------------------------- rendering

def _text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _grid(n):
    return {1: (1, 1), 2: (2, 1), 3: (3, 1), 4: (2, 2)}[n]


def _place_images(slide, imgs, x0, y0, w, h, gap=Inches(0.25)):
    cols, rows = _grid(len(imgs))
    cell_w = (w - gap * (cols - 1)) / cols
    cell_h = (h - gap * (rows - 1)) / rows
    for k, (data, iw, ih) in enumerate(imgs):
        r, c = divmod(k, cols)
        scale = min(cell_w / iw, cell_h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        cx = int(x0 + c * (cell_w + gap) + (cell_w - pw) / 2)
        cy = int(y0 + r * (cell_h + gap) + (cell_h - ph) / 2)
        slide.shapes.add_picture(io.BytesIO(data), cx, cy, pw, ph)


def build_deck(plan: dict, images: list) -> bytes:
    """images[i] = (jpeg_bytes, width, height), indexed the same way as the plan."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    content_w = SLIDE_W - 2 * MARGIN

    # Title slide
    s = prs.slides.add_slide(blank)
    _background(s)
    bar = s.shapes.add_shape(1, MARGIN, Inches(3.05), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _text(s, MARGIN, Inches(1.4), content_w, Inches(1.6), plan["title"], 44, INK,
          bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if plan["subtitle"]:
        _text(s, MARGIN, Inches(3.35), content_w, Inches(1.0), plan["subtitle"], 22, MUTED)

    # Content slides
    for spec in plan["slides"]:
        s = prs.slides.add_slide(blank)
        _background(s)
        _text(s, MARGIN, Inches(0.4), content_w, Inches(0.9), spec["title"], 30, INK,
              bold=True, anchor=MSO_ANCHOR.MIDDLE)

        has_caption = bool(spec["caption"])
        top = Inches(1.45)
        bottom = SLIDE_H - (Inches(1.35) if has_caption else MARGIN)
        imgs = [images[i] for i in spec["images"] if images[i] is not None]
        if imgs:
            _place_images(s, imgs, MARGIN, top, content_w, bottom - top)
        if has_caption:
            _text(s, MARGIN, SLIDE_H - Inches(1.15), content_w, Inches(0.8), spec["caption"],
                  18, MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if spec["notes"]:
            s.notes_slide.notes_text_frame.text = spec["notes"]

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# Quality ladder: try the best first, step down until the file fits Discord's upload cap.
QUALITY_STEPS = [(1920, 85), (1400, 78), (1024, 70), (720, 60)]


def build_deck_under(plan: dict, raw_images: list[bytes], limit_bytes: int) -> bytes | None:
    for max_side, q in QUALITY_STEPS:
        imgs = [normalize_image(b, max_side, q) for b in raw_images]
        data = build_deck(plan, imgs)
        if len(data) <= limit_bytes:
            return data
    return None
